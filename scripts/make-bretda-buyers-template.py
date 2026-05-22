#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera planilha modelo para vendedores da Bretda preencherem compradores reais."""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.expanduser("~"), "Downloads", "Bretda-Compradores-MODELO.xlsx")

wb = Workbook()

# ---------- Cores / estilos ----------
DARK = "1F2A33"      # header escuro
AMBER = "E8932E"
LIGHT = "F4F1EC"
GRAY = "9AA0A6"
white = Font(color="FFFFFF", bold=True, size=11)
hdr_fill = PatternFill("solid", fgColor=DARK)
ex_font = Font(color=GRAY, italic=True, size=10)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(vertical="center", wrap_text=True)
center = Alignment(horizontal="center", vertical="center")

# ============================================================
# SHEET 1 — Compradores
# ============================================================
ws = wb.active
ws.title = "Compradores"

# Linha 1: aviso topo
ws.merge_cells("A1:H1")
c = ws["A1"]
c.value = ("⚠️  PREENCHA SÓ QUEM COMPROU DE VERDADE (não leads, não orçamentos abertos).  "
           "Telefone no formato +55 + DDD + número.  Apague as 2 linhas de EXEMPLO antes de devolver.")
c.font = Font(bold=True, color="7A4A00", size=10)
c.fill = PatternFill("solid", fgColor="FCEFD8")
c.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 38

# Linha 2: cabeçalhos
headers = ["Nome", "Sobrenome", "Telefone (WhatsApp)", "Email",
           "Cidade", "UF", "Mesa comprada", "Data da compra (mm/aaaa)"]
widths = [16, 16, 22, 28, 18, 6, 22, 22]
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    cell = ws.cell(row=2, column=i, value=h)
    cell.font = white
    cell.fill = hdr_fill
    cell.alignment = center
    cell.border = border
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[2].height = 28

# Linhas 3-4: exemplos
examples = [
    ["João", "Silva", "+5547999998888", "joao.silva@email.com", "Blumenau", "SC", "Aurora", "03/2026"],
    ["Maria", "Costa", "+5511988887777", "maria@email.com", "São Paulo", "SP", "Citrino", "11/2025"],
]
for r, row in enumerate(examples, start=3):
    for i, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=i, value=val)
        cell.font = ex_font
        cell.border = border
        cell.alignment = wrap
    ws.cell(row=r, column=1).value = "EXEMPLO — apague › " + row[0]

# Linhas em branco prontas pra preencher (com bordas)
for r in range(5, 60):
    for i in range(1, 9):
        ws.cell(row=r, column=i).border = border

ws.freeze_panes = "A3"   # mantém aviso + cabeçalho visíveis

# ============================================================
# SHEET 2 — Instruções
# ============================================================
ins = wb.create_sheet("Leia antes de preencher")
ins.column_dimensions["A"].width = 100

ins["A1"] = "Como preencher a lista de compradores — Bretda"
ins["A1"].font = Font(bold=True, size=14, color=DARK)
ins.row_dimensions[1].height = 26

linhas = [
    "",
    "Objetivo: montar a lista de quem JÁ COMPROU uma mesa Bretda. Essa lista vira o público que o",
    "anúncio vai usar pra achar mais gente parecida com comprador de verdade.",
    "",
    "REGRAS (importante):",
    "1.  Liste SÓ clientes que COMPRARAM. Nada de lead, orçamento aberto, 'tava interessado' ou curioso.",
    "2.  Uma pessoa por linha. Se a pessoa comprou mais de uma vez, deixe só uma linha.",
    "3.  Telefone (o mais importante): formato  +55 + DDD + número, sem espaço, traço ou parênteses.",
    "         ✅ certo:  +5547999998888           ❌ errado:  (47) 99999-8888",
    "4.  Email: se tiver, preencha — melhora muito o resultado.",
    "5.  Não tem certeza de um dado? Deixe em branco. Melhor vazio do que errado.",
    "6.  As colunas 'Mesa comprada' e 'Data da compra' são só pra conferência interna — não vão pro anúncio,",
    "         mas ajudam a garantir que é comprador real. Preencha se souber.",
    "7.  Apague as 2 linhas de EXEMPLO antes de devolver.",
    "8.  Pode preencher no Excel ou subir no Google Sheets, tanto faz.",
    "",
    "Quanto mais compradores na lista, melhor o resultado (mínimo ~100).",
    "",
    "Quando terminar, devolva este arquivo preenchido.",
]
for i, t in enumerate(linhas, start=2):
    cell = ins.cell(row=i, column=1, value=t)
    if t.startswith("REGRAS"):
        cell.font = Font(bold=True, size=11, color=DARK)
    elif t and t[0].isdigit():
        cell.font = Font(size=11)
    else:
        cell.font = Font(size=11, color="3C3C3C")

wb.save(OUT)
print("OK ->", OUT)
