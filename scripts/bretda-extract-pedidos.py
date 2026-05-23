# -*- coding: utf-8 -*-
"""Extract contact + delivery data from Bretda orders/proposals (PDF + DOCX)."""
import os, re, sys, json, glob

SRC = sys.argv[1]

def read_pdf(p):
    import fitz
    t = []
    with fitz.open(p) as doc:
        for page in doc:
            t.append(page.get_text())
    return "\n".join(t)

def read_docx(p):
    import docx
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)

EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
# Brazilian phone: optional +55, DDD, 8-9 digits, with separators
PHONE = re.compile(r"(?:\+?55\s*)?\(?\d{2}\)?[\s.\-]?9?\d{4}[\s.\-]?\d{4}")
CEP   = re.compile(r"\b\d{5}-?\d{3}\b")
CPF   = re.compile(r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b")
CNPJ  = re.compile(r"\b\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}\b")
UF    = re.compile(r"\b(AC|AL|AP|AM|BA|CE|DF|ES|GO|MA|MT|MS|MG|PA|PB|PR|PE|PI|RJ|RN|RS|RO|RR|SC|SP|SE|TO)\b")

def clean_phones(text):
    out = set()
    for m in PHONE.findall(text):
        digits = re.sub(r"\D", "", m)
        if digits.startswith("55") and len(digits) > 11:
            digits = digits[2:]
        if 10 <= len(digits) <= 11:
            out.add(digits)
    return sorted(out)

results = []
files = sorted(glob.glob(os.path.join(SRC, "*")))
for f in files:
    name = os.path.basename(f)
    low = name.lower()
    if low.endswith(".pdf"):
        try: text = read_pdf(f)
        except Exception as e: text = f"<<ERR {e}>>"
    elif low.endswith(".docx"):
        try: text = read_docx(f)
        except Exception as e: text = f"<<ERR {e}>>"
    else:
        continue
    kind = "OC" if name.upper().startswith("OC") else ("PROPOSTA" if "PROPOSTA" in name.upper() else "OUTRO")
    emails = sorted(set(EMAIL.findall(text)))
    phones = clean_phones(text)
    ceps   = sorted(set(CEP.findall(text)))
    cpfs   = sorted(set(CPF.findall(text)))
    cnpjs  = sorted(set(CNPJ.findall(text)))
    ufs    = sorted(set(UF.findall(text)))
    results.append({
        "file": name, "kind": kind,
        "emails": emails, "phones": phones,
        "ceps": ceps, "cpfs": cpfs, "cnpjs": cnpjs, "ufs": ufs,
        "chars": len(text),
        "text": text,
    })

# Save full dump
with open(os.path.join(SRC, "_dump.json"), "w", encoding="utf-8") as fh:
    json.dump(results, fh, ensure_ascii=False, indent=1)

# Print summary table
print(f"{'KIND':9} {'EMAIL':5} {'TEL':3} {'CEP':3} {'CPF':3} {'CNPJ':4} FILE")
for r in results:
    print(f"{r['kind']:9} {len(r['emails']):<5} {len(r['phones']):<3} {len(r['ceps']):<3} {len(r['cpfs']):<3} {len(r['cnpjs']):<4} {r['file']}")

oc = [r for r in results if r['kind']=='OC']
pp = [r for r in results if r['kind']=='PROPOSTA']
print(f"\n== TOTAIS ==")
print(f"arquivos: {len(results)} | OC(pedidos): {len(oc)} | PROPOSTA(orcamentos): {len(pp)}")
print(f"OC com email: {sum(1 for r in oc if r['emails'])} | OC com tel: {sum(1 for r in oc if r['phones'])}")
print(f"PP com email: {sum(1 for r in pp if r['emails'])} | PP com tel: {sum(1 for r in pp if r['phones'])}")
allmail = set(); allphone=set()
for r in results:
    allmail.update(r['emails']); allphone.update(r['phones'])
print(f"emails unicos (todos): {len(allmail)} | telefones unicos (todos): {len(allphone)}")
