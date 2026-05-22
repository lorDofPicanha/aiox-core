#!/usr/bin/env python3
"""Build PCMSO .docx — corrected text with VISIBLE edit markers.

Same as v2 but each edit is rendered with:
- INSERTIONS: green text
- REMOVALS: red strikethrough
- REPLACEMENTS: red strikethrough (old) + green (new)
"""
import sys
import re
from pathlib import Path
import fitz
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

if len(sys.argv) < 2:
    print("Usage: pcmso-build-docx-v3.py <pdf_path>"); sys.exit(1)

pdf_path = Path(sys.argv[1])
out_path = Path(f"C:/Users/kingp/Downloads/PCMSO-Corrigidos/{pdf_path.stem} - CORRIGIDO.docx")

# Markers — chosen to avoid any chance of natural collision
DEL_OPEN, DEL_CLOSE = "⟦DEL⟧", "⟦/DEL⟧"
INS_OPEN, INS_CLOSE = "⟦INS⟧", "⟦/INS⟧"

GREEN = RGBColor(0x00, 0x80, 0x00)
RED = RGBColor(0xC0, 0x00, 0x00)


def categorize(content):
    c = (content or "").strip()
    if not c:
        return ("SKIP", "")
    cl = c.lower()
    m = re.match(r'^(?:retirar|tirar|remover|excluir|deletar)\s+(.+)$', c, re.IGNORECASE | re.DOTALL)
    if m:
        body = m.group(1)
        m2 = re.search(r'(?:palavra|termo|express[ãa]o|frase)\s+["\']?([^"\']+?)["\']?\s*\.?\s*$', body, re.IGNORECASE)
        if m2:
            return ("REMOVE", m2.group(1).strip())
        return ("REMOVE", body.strip(' .,;:!?"\''))
    m = re.match(r'^(?:acrescentar|adicionar|incluir|colocar|inserir|por)\s+(.+)$', c, re.IGNORECASE | re.DOTALL)
    if m:
        return ("INSERT", m.group(1).strip())
    if cl.endswith('?') or re.match(r'^(verificar|esclarecer|pretende|conferir)\b', cl):
        return ("SKIP", "")
    return ("REPLACE", c)


def get_text_in_rect(page, rect):
    words = page.get_text("words")
    matched = []
    for w in words:
        wrect = fitz.Rect(w[:4])
        if rect.intersects(wrect):
            inter = rect & wrect
            if inter.get_area() / max(wrect.get_area(), 1) > 0.3:
                matched.append(w[4])
    return " ".join(matched).strip()


def extract_marked(page, annot):
    verts = annot.vertices
    if verts and len(verts) >= 4 and len(verts) % 4 == 0:
        spans = []
        for i in range(0, len(verts), 4):
            quad = verts[i:i+4]
            xs = [p[0] for p in quad]; ys = [p[1] for p in quad]
            qrect = fitz.Rect(min(xs), min(ys), max(xs), max(ys))
            spans.append(get_text_in_rect(page, qrect))
        return " ".join(s for s in spans if s).strip()
    return get_text_in_rect(page, annot.rect)


def add_paragraph_with_markers(doc, text):
    """Parse text with ⟦DEL⟧/⟦INS⟧ markers and add as paragraph with formatted runs."""
    p = doc.add_paragraph()
    pattern = re.compile(rf'({re.escape(DEL_OPEN)}.*?{re.escape(DEL_CLOSE)}|{re.escape(INS_OPEN)}.*?{re.escape(INS_CLOSE)})', re.DOTALL)
    pos = 0
    for m in pattern.finditer(text):
        # Plain text before marker
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith(DEL_OPEN):
            inner = token[len(DEL_OPEN):-len(DEL_CLOSE)]
            r = p.add_run(inner)
            r.font.color.rgb = RED
            r.font.strike = True
        elif token.startswith(INS_OPEN):
            inner = token[len(INS_OPEN):-len(INS_CLOSE)]
            r = p.add_run(inner)
            r.font.color.rgb = GREEN
            r.bold = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])
    return p


# Process PDF
src = fitz.open(pdf_path)
out_doc = Document()
title = out_doc.add_heading(pdf_path.stem.replace(" (1)", ""), level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Legend
p_leg = out_doc.add_paragraph()
p_leg.add_run('Legenda: ').bold = True
r_ins = p_leg.add_run('verde negrito')
r_ins.font.color.rgb = GREEN; r_ins.bold = True
p_leg.add_run(' = inserido · ')
r_del = p_leg.add_run('vermelho riscado')
r_del.font.color.rgb = RED; r_del.font.strike = True
p_leg.add_run(' = removido/substituído')
p_leg.runs[0].font.size = Pt(9)
for r in p_leg.runs:
    r.font.size = Pt(9)

style = out_doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

applied = 0
skipped = 0

for page_num, page in enumerate(src, start=1):
    annots_data = []
    for annot in page.annots() or []:
        content = (annot.info.get("content", "") or "").strip()
        if not content:
            continue
        marked = extract_marked(page, annot)
        action, payload = categorize(content)
        annots_data.append({
            "marked": marked, "content": content,
            "action": action, "payload": payload,
            "rect": annot.rect,
        })

    page_text = page.get_text("text")

    # Sort top-down
    annots_data.sort(key=lambda a: (a["rect"].y0, a["rect"].x0))

    for a in annots_data:
        if a["action"] == "SKIP" or not a["marked"]:
            skipped += 1
            continue

        marked = a["marked"]
        norm_marked = re.sub(r'\s+', ' ', marked).strip()
        words = [re.escape(w) for w in norm_marked.split() if len(w) > 2][:6]
        if not words and a["action"] != "REMOVE":
            skipped += 1
            continue

        if a["action"] == "REMOVE":
            word = a["payload"]
            pat = re.compile(r'\b(' + re.escape(word) + r')\b', re.IGNORECASE)
            new_text, n = pat.subn(f"{DEL_OPEN}\\1{DEL_CLOSE}", page_text, count=1)
            if n > 0:
                page_text = new_text
                applied += 1
            else:
                skipped += 1
        elif a["action"] == "INSERT":
            pattern = re.compile(r'\s+'.join(words), re.IGNORECASE | re.DOTALL)
            m = pattern.search(page_text)
            if m:
                page_text = page_text[:m.end()] + f" {INS_OPEN}{a['payload']}{INS_CLOSE}" + page_text[m.end():]
                applied += 1
            else:
                skipped += 1
        elif a["action"] == "REPLACE":
            pattern = re.compile(r'\s+'.join(words), re.IGNORECASE | re.DOTALL)
            m = pattern.search(page_text)
            if m:
                page_text = (
                    page_text[:m.start()]
                    + f"{DEL_OPEN}{m.group(0)}{DEL_CLOSE}"
                    + f"{INS_OPEN}{a['payload']}{INS_CLOSE}"
                    + page_text[m.end():]
                )
                applied += 1
            else:
                skipped += 1

    # Write the page text as paragraphs, parsing markers
    for para in re.split(r'\n\s*\n', page_text):
        para = para.strip()
        if not para:
            continue
        clean = re.sub(r'\s*\n\s*', ' ', para)
        clean = re.sub(r'  +', ' ', clean).strip()
        if clean:
            add_paragraph_with_markers(out_doc, clean)

src.close()
out_doc.save(out_path)
print(f"Generated: {out_path}")
print(f"Applied: {applied} | Skipped: {skipped}")
